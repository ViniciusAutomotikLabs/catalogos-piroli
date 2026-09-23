#!/usr/bin/env python3
"""Importa planilhas de estoque SS → Redis gpasi:peca:{codigo}.

Uso (VPS / container com redis):
  python gpasi_import_stock_xlsx.py \\
    --xlsx-28 /opt/gpasi-sync/imports/00028GUS.xlsx \\
    --xlsx-35 /opt/gpasi-sync/imports/00035GUS.xlsx

Regras:
- 00035 é fonte principal (estoque 0001/0003/0004/0005 + preços da planilha).
- 00028 completa marca / grupo_descricao e cobre códigos só nela.
- Só escreve campos de estoque (+ preços da 35) — NÃO apaga aplicacao/tok/enrich.
- Códigos sem hash: cria hash mínimo (codigo/descricao/marca/fab + estoque).
- Contadores e índices gpasi:tok:* / gpasi:aplic:* não são alterados.
"""

from __future__ import annotations

import argparse
import os
import sys
import time
from datetime import datetime, timezone
from typing import Any

try:
    import openpyxl
except ImportError:
    print("precisa openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(2)

try:
    import redis
except ImportError:
    print("precisa redis: pip install redis", file=sys.stderr)
    sys.exit(2)


def parse_num(v: Any) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _s(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def load_28(path: str) -> dict[str, dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = [_s(c) for c in next(rows)]
    # achar sheet certa se a primeira for vazia
    if "Codigo_Interno" not in hdr:
        for name in wb.sheetnames:
            ws = wb[name]
            rows = ws.iter_rows(values_only=True)
            hdr = [_s(c) for c in next(rows)]
            if "Codigo_Interno" in hdr:
                break
    idx = {h: i for i, h in enumerate(hdr)}
    out: dict[str, dict[str, Any]] = {}
    for row in rows:
        code = _s(row[idx["Codigo_Interno"]])
        if not code:
            continue
        out[code] = {
            "marca": _s(row[idx.get("Marca", 1)]),
            "descricao": _s(row[idx.get("Descricao_do_produto", 2)]),
            "codigofabricante": _s(row[idx.get("Codigo_Fabricante", 3)]),
            "estoque_disponivel": parse_num(row[idx["Estoque_Disponivel"]]) or 0.0,
            "estoque_0001": parse_num(row[idx["Estoque_0001"]]) or 0.0,
            "estoque_0004": parse_num(row[idx["Estoque_0004"]]) or 0.0,
            "estoque_0005": parse_num(row[idx["Estoque_0005"]]) or 0.0,
            "grupo_descricao": _s(row[idx.get("Grupo_Descricao", 8)]),
        }
    wb.close()
    return out


def load_35(path: str) -> dict[str, dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = None
    rows = None
    hdr: list[str] = []
    for name in wb.sheetnames:
        cand = wb[name]
        it = cand.iter_rows(values_only=True)
        h = [_s(c) for c in next(it)]
        if "Codigo_Interno" in h and "Estoque_0003" in h:
            ws, rows, hdr = cand, it, h
            break
    if rows is None:
        raise RuntimeError(f"sheet 00035 não encontrada em {path}")
    idx = {h: i for i, h in enumerate(hdr)}
    out: dict[str, dict[str, Any]] = {}
    for row in rows:
        code = _s(row[idx["Codigo_Interno"]])
        if not code:
            continue
        e1 = parse_num(row[idx["Estoque_0001"]]) or 0.0
        e3 = parse_num(row[idx["Estoque_0003"]]) or 0.0
        e4 = parse_num(row[idx["Estoque_0004"]]) or 0.0
        e5 = parse_num(row[idx["Estoque_0005"]]) or 0.0
        out[code] = {
            "codigofabricante": _s(row[idx["Codigo_Fabricante"]]),
            "descricao": _s(row[idx["Descricao_do_produto"]]),
            "preco": parse_num(row[idx["Preco_a_Prazo"]]),
            "preco_promoc_varejo": parse_num(row[idx["Preco_Promoc_Varejo"]]),
            "precoatacado": parse_num(row[idx["Preco_de_Atacado"]]),
            "estoque_0001": e1,
            "estoque_0003": e3,
            "estoque_0004": e4,
            "estoque_0005": e5,
            "estoque": e1 + e3 + e4 + e5,
        }
    wb.close()
    return out


def merge(a28: dict[str, dict[str, Any]], a35: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """35 manda no estoque/preço; 28 completa marca/grupo e códigos exclusivos."""
    out: dict[str, dict[str, Any]] = {}
    for code, v in a35.items():
        m: dict[str, Any] = {
            "estoque": v["estoque"],
            "estoque_0001": v["estoque_0001"],
            "estoque_0003": v["estoque_0003"],
            "estoque_0004": v["estoque_0004"],
            "estoque_0005": v["estoque_0005"],
            "descricao": v["descricao"],
            "codigofabricante": v["codigofabricante"],
            "preco": v["preco"],
            "preco_promoc_varejo": v["preco_promoc_varejo"],
            "precoatacado": v["precoatacado"],
            "stock_source": "xlsx:00035GUS",
        }
        if code in a28:
            if a28[code].get("marca"):
                m["marca"] = a28[code]["marca"]
            if a28[code].get("grupo_descricao"):
                m["grupo_descricao"] = a28[code]["grupo_descricao"]
            if not m["descricao"] and a28[code].get("descricao"):
                m["descricao"] = a28[code]["descricao"]
        out[code] = m

    for code, v in a28.items():
        if code in out:
            continue
        out[code] = {
            "estoque": v["estoque_disponivel"],
            "estoque_0001": v["estoque_0001"],
            "estoque_0003": 0.0,
            "estoque_0004": v["estoque_0004"],
            "estoque_0005": v["estoque_0005"],
            "descricao": v["descricao"],
            "codigofabricante": v["codigofabricante"],
            "marca": v.get("marca") or "",
            "grupo_descricao": v.get("grupo_descricao") or "",
            "preco": None,
            "preco_promoc_varejo": None,
            "precoatacado": None,
            "stock_source": "xlsx:00028GUS",
        }
    return out


def fmt_qty(n: float) -> str:
    # inteiro se .0, senão até 3 casas
    if abs(n - round(n)) < 1e-9:
        return str(int(round(n)))
    return f"{n:.3f}".rstrip("0").rstrip(".")


def import_redis(r: redis.Redis, merged: dict[str, dict[str, Any]], update_prices: bool) -> dict[str, int]:
    now = datetime.now(timezone.utc).isoformat()
    pipe = r.pipeline(transaction=False)
    updated = 0
    created = 0
    with_stock = 0
    chunk = 500
    codes = list(merged.keys())

    # existência em lotes
    exists_map: dict[str, bool] = {}
    for i in range(0, len(codes), chunk):
        batch = codes[i : i + chunk]
        for c in batch:
            pipe.exists(f"gpasi:peca:{c}")
        flags = pipe.execute()
        for c, flag in zip(batch, flags):
            exists_map[c] = bool(flag)

    pipe = r.pipeline(transaction=False)
    pending = 0
    for code, v in merged.items():
        key = f"gpasi:peca:{code}"
        mapping: dict[str, str] = {
            "estoque": fmt_qty(float(v["estoque"])),
            "estoque_0001": fmt_qty(float(v["estoque_0001"])),
            "estoque_0003": fmt_qty(float(v["estoque_0003"])),
            "estoque_0004": fmt_qty(float(v["estoque_0004"])),
            "estoque_0005": fmt_qty(float(v["estoque_0005"])),
            "stock_updated_at": now,
            "stock_source": str(v.get("stock_source") or "xlsx"),
        }
        if float(v["estoque"]) > 0:
            with_stock += 1

        if not exists_map.get(code):
            mapping["codigo"] = code
            if v.get("descricao"):
                mapping["descricao"] = str(v["descricao"])[:500]
            if v.get("codigofabricante"):
                mapping["codigofabricante"] = str(v["codigofabricante"])[:80]
            if v.get("marca"):
                mapping["marca"] = str(v["marca"])[:120]
            if v.get("grupo_descricao"):
                mapping["grupo_descricao"] = str(v["grupo_descricao"])[:120]
            created += 1
        else:
            # em hashes existentes: não sobrescreve descricao/marca/fab do ERP
            updated += 1

        if update_prices:
            if v.get("preco") is not None:
                mapping["preco"] = f"{float(v['preco']):.2f}"
            if v.get("precoatacado") is not None:
                mapping["precoatacado"] = f"{float(v['precoatacado']):.2f}"
            if v.get("preco_promoc_varejo") is not None:
                mapping["preco_promoc_varejo"] = f"{float(v['preco_promoc_varejo']):.2f}"
            mapping["price_sheet_updated_at"] = now

        pipe.hset(key, mapping=mapping)
        pending += 1
        if pending >= chunk:
            pipe.execute()
            pipe = r.pipeline(transaction=False)
            pending = 0

    if pending:
        pipe.execute()

    r.set("gpasi:meta:stock_updated_at", now)
    r.set("gpasi:meta:stock_count", str(len(merged)))
    r.set("gpasi:meta:stock_with_qty", str(with_stock))
    r.set("gpasi:meta:stock_source", "xlsx:00028GUS+00035GUS")
    return {
        "total": len(merged),
        "updated": updated,
        "created": created,
        "with_stock": with_stock,
    }


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--xlsx-28", required=True)
    p.add_argument("--xlsx-35", required=True)
    p.add_argument("--update-prices", action="store_true",
                   help="também grava preco/precoatacado da planilha 35")
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()

    t0 = time.perf_counter()
    print(f"loading 28: {args.xlsx_28}", flush=True)
    a28 = load_28(args.xlsx_28)
    print(f"  {len(a28):,} codes", flush=True)
    print(f"loading 35: {args.xlsx_35}", flush=True)
    a35 = load_35(args.xlsx_35)
    print(f"  {len(a35):,} codes", flush=True)
    merged = merge(a28, a35)
    gt0 = sum(1 for v in merged.values() if float(v["estoque"]) > 0)
    print(f"merged {len(merged):,} codes, with_stock>0={gt0:,}", flush=True)

    if args.dry_run:
        sample = [(c, v) for c, v in merged.items() if float(v["estoque"]) > 0][:5]
        for c, v in sample:
            print(" sample", c, v["estoque"], v.get("descricao", "")[:40])
        return 0

    host = os.environ.get("REDIS_HOST", "127.0.0.1")
    port = int(os.environ.get("REDIS_PORT", "6379"))
    password = os.environ.get("REDIS_PASSWORD") or None
    r = redis.Redis(host=host, port=port, password=password, decode_responses=True)
    r.ping()
    print(f"redis={host}:{port} importing…", flush=True)
    stats = import_redis(r, merged, update_prices=args.update_prices)
    print(
        f"done in {time.perf_counter()-t0:.1f}s: "
        f"total={stats['total']:,} updated={stats['updated']:,} "
        f"created={stats['created']:,} with_stock={stats['with_stock']:,}",
        flush=True,
    )
    # sanity
    for probe in ("000781", "000062", "012729"):
        if r.exists(f"gpasi:peca:{probe}"):
            print(
                "probe",
                probe,
                r.hmget(
                    f"gpasi:peca:{probe}",
                    "estoque",
                    "estoque_0001",
                    "estoque_0003",
                    "estoque_0004",
                    "estoque_0005",
                    "stock_source",
                ),
            )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
